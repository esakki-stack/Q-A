"""
Document Extraction Service
Handles extraction from URLs, PDFs, Word documents, and plain text
"""

import asyncio
import logging
from typing import Optional, Tuple
from datetime import datetime
from tenacity import retry, stop_after_attempt, wait_exponential

try:
    from newspaper import Article
    NEWSPAPER_AVAILABLE = True
except ImportError:
    NEWSPAPER_AVAILABLE = False

try:
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

import requests
from app.schemas import ArticleMetadata

logger = logging.getLogger(__name__)


class ExtractionService:
    """
    Unified service for extracting text and metadata from multiple sources
    Supports: URLs (via newspaper3k), PDFs (via pdfplumber), Word docs, plain text
    """
    
    @staticmethod
    @retry(
        stop=stop_after_attempt(3),
        wait=wait_exponential(multiplier=1, min=2, max=10),
        reraise=True
    )
    async def extract_from_url(url: str, timeout: int = 15) -> Tuple[str, ArticleMetadata]:
        """
        Extract article text and metadata from URL
        
        Args:
            url: URL of the article
            timeout: Request timeout in seconds
            
        Returns:
            Tuple of (article_text, metadata)
        """
        if not NEWSPAPER_AVAILABLE:
            raise ValueError("newspaper3k not installed. Install with: pip install newspaper3k")
        
        logger.info(f"📰 Extracting from URL: {url}")
        
        try:
            # Run in thread pool to avoid blocking
            loop = asyncio.get_event_loop()
            
            article = Article(str(url), timeout=timeout)
            await loop.run_in_executor(None, article.download)
            await loop.run_in_executor(None, article.parse)
            
            # Extract metadata
            metadata = ArticleMetadata(
                title=article.title,
                author=article.authors[0] if article.authors else None,
                publish_date=article.publish_date,
                source_url=str(url),
                word_count=len(article.text.split()),
                language=article.meta_lang or "en",
                categories=article.categories if hasattr(article, 'categories') else []
            )
            
            # Calculate reading time (average 200 words per minute)
            metadata.reading_time_minutes = metadata.word_count / 200
            
            logger.info(f"✅ Extracted {metadata.word_count} words from {url}")
            
            return article.text, metadata
            
        except Exception as e:
            logger.error(f"❌ Error extracting from URL {url}: {e}")
            raise
    
    
    @staticmethod
    async def extract_from_pdf(file_path: str) -> Tuple[str, ArticleMetadata]:
        """
        Extract text from PDF file
        
        Args:
            file_path: Path to PDF file
            
        Returns:
            Tuple of (text, metadata)
        """
        if not PDF_AVAILABLE:
            raise ValueError("pdfplumber not installed. Install with: pip install pdfplumber")
        
        logger.info(f"📄 Extracting from PDF: {file_path}")
        
        try:
            text_parts = []
            
            with pdfplumber.open(file_path) as pdf:
                num_pages = len(pdf.pages)
                
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            
            full_text = "\n".join(text_parts)
            word_count = len(full_text.split())
            
            metadata = ArticleMetadata(
                title=None,
                word_count=word_count,
                reading_time_minutes=word_count / 200,
                language="en",
            )
            
            logger.info(f"✅ Extracted {word_count} words from PDF ({num_pages} pages)")
            
            return full_text, metadata
            
        except Exception as e:
            logger.error(f"❌ Error extracting PDF: {e}")
            raise
    
    
    @staticmethod
    async def extract_from_docx(file_path: str) -> Tuple[str, ArticleMetadata]:
        """
        Extract text from Word document
        
        Args:
            file_path: Path to DOCX file
            
        Returns:
            Tuple of (text, metadata)
        """
        if not DOCX_AVAILABLE:
            raise ValueError("python-docx not installed. Install with: pip install python-docx")
        
        logger.info(f"📝 Extracting from DOCX: {file_path}")
        
        try:
            doc = Document(file_path)
            
            # Extract all paragraphs
            text_parts = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            full_text = "\n".join(text_parts)
            word_count = len(full_text.split())
            
            # Try to get title from first heading
            title = None
            for paragraph in doc.paragraphs:
                if paragraph.style.name.startswith("Heading"):
                    title = paragraph.text
                    break
            
            metadata = ArticleMetadata(
                title=title,
                word_count=word_count,
                reading_time_minutes=word_count / 200,
                language="en",
            )
            
            logger.info(f"✅ Extracted {word_count} words from DOCX")
            
            return full_text, metadata
            
        except Exception as e:
            logger.error(f"❌ Error extracting DOCX: {e}")
            raise
    
    
    @staticmethod
    def extract_from_text(text: str, source_url: Optional[str] = None) -> Tuple[str, ArticleMetadata]:
        """
        Process plain text (already provided)
        
        Args:
            text: Plain text content
            source_url: Optional source URL
            
        Returns:
            Tuple of (text, metadata)
        """
        logger.info(f"📋 Processing plain text ({len(text)} chars)")
        
        word_count = len(text.split())
        
        metadata = ArticleMetadata(
            source_url=source_url,
            word_count=word_count,
            reading_time_minutes=word_count / 200,
            language="en",
        )
        
        logger.info(f"✅ Processed {word_count} words from text")
        
        return text, metadata
    
    
    @staticmethod
    async def extract(
        article_text: Optional[str] = None,
        article_url: Optional[str] = None,
        file_path: Optional[str] = None,
        file_type: Optional[str] = None,
    ) -> Tuple[str, ArticleMetadata]:
        """
        Unified extraction method - routes to appropriate extractor
        
        Args:
            article_text: Plain text content
            article_url: URL to extract from
            file_path: Path to file
            file_type: Type of file (pdf, docx, etc.)
            
        Returns:
            Tuple of (extracted_text, metadata)
            
        Raises:
            ValueError: If no valid input is provided
        """
        # Priority: text > url > file
        if article_text and article_text.strip():
            logger.info("Using article text")
            return ExtractionService.extract_from_text(article_text, source_url=article_url)
        
        elif article_url:
            logger.info("Using URL extraction")
            return await ExtractionService.extract_from_url(article_url)
        
        elif file_path:
            logger.info(f"Using file extraction ({file_type})")
            
            if not file_type:
                # Guess from extension
                file_type = file_path.split(".")[-1].lower()
            
            if file_type == "pdf":
                return await ExtractionService.extract_from_pdf(file_path)
            elif file_type in ("docx", "doc"):
                return await ExtractionService.extract_from_docx(file_path)
            else:
                # Try text
                with open(file_path, "r", encoding="utf-8") as f:
                    text = f.read()
                return ExtractionService.extract_from_text(text, source_url=None)
        
        else:
            raise ValueError("Must provide either article_text, article_url, or file_path")
